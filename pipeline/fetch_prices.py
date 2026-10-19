"""Prezzi settimanali (EU Weekly Oil Bulletin) per l'Italia, accise e IVA, più il prezzo UE al netto delle tasse
(usato per la stagionalità del controfattuale, come nel tracker UE)."""
import datetime as dt

import openpyxl

from common import CACHE, WOB_URL, http_get

FUELS = {"benzina": "euro95", "gasolio": "diesel"}


def download(force: bool) -> str:
    CACHE.mkdir(parents=True, exist_ok=True)
    path = CACHE / "wob_history.xlsx"
    if force or not path.exists():
        path.write_bytes(http_get(WOB_URL, timeout=300))
    return path


def _prices(ws, kind: str, area: str) -> dict:
    rows = ws.iter_rows(values_only=True)
    col = {h: i for i, h in enumerate(next(rows)) if isinstance(h, str)}
    out = {f: {} for f in FUELS}
    for r in rows:
        if not isinstance(r[0], dt.datetime):
            continue
        for f, suffix in FUELS.items():
            v = r[col[f"{area}_price_{kind}_{suffix}"]]
            if isinstance(v, (int, float)) and v > 0:
                out[f][r[0].date()] = v / 1000.0
    return out


def _history(ws, area: str) -> dict:
    """{fuel: {since_date: value}} for Italy from a tax-history sheet (VAT % or excise €/1000 L)."""
    out, cur = {f: {} for f in FUELS}, None
    for r in list(ws.iter_rows(values_only=True))[4:]:
        if isinstance(r[0], str) and r[0].endswith("_"):
            cur = r[0][:-1]
        if cur != area or not isinstance(r[1], dt.datetime):
            continue
        for i, f in enumerate(FUELS):
            if r[2 + i] is not None:
                out[f][r[1].date()] = float(r[2 + i])
    return out


def load(force: bool = True) -> dict:
    wb = openpyxl.load_workbook(download(force), read_only=True, data_only=True)
    data = {
        "retail": _prices(wb["Prices with taxes"], "with_tax", "IT"),
        "pretax": _prices(wb["Prices wo taxes"], "wo_tax", "IT"),
        "eu_pretax": _prices(wb["Prices wo taxes"], "wo_tax", "EU"),
        "excise": {f: {d: v / 1000.0 for d, v in h.items()} for f, h in _history(wb["Excise duties"], "IT").items()},
        "vat": {f: {d: v / 100.0 for d, v in h.items()} for f, h in _history(wb["VAT"], "IT").items()},
    }
    wb.close()
    return data


if __name__ == "__main__":
    d = load(force=True)
    last = max(d["retail"]["gasolio"])
    print("latest bulletin", last, {f: round(d["retail"][f][last], 3) for f in FUELS})
    print("diesel excise changes 2026:", {k.isoformat(): v for k, v in sorted(d["excise"]["gasolio"].items()) if k.year >= 2026})
    print("petrol excise changes 2026:", {k.isoformat(): v for k, v in sorted(d["excise"]["benzina"].items()) if k.year >= 2026})
